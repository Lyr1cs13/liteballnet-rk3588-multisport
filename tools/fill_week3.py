from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

SOURCE = Path(r"C:\\Users\\97689\\AppData\\Local\\Temp\\week3_blank.docx")
OUTPUT = Path(r"C:\\Users\\97689\\AppData\\Local\\Temp\\week3_liteballnet.docx")
FRAMES = [
    Path(r"C:\\Users\\97689\\AppData\\Local\\Temp\\tennis_reference_fp.png"),
    Path(r"C:\\Users\\97689\\AppData\\Local\\Temp\\table_tennis_reference_fp.png"),
    Path(r"C:\\Users\\97689\\AppData\\Local\\Temp\\badminton_reference_fp.png"),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_run(run, bold=False, size=10.5):
    run.font.name = '宋体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(size)
    run.bold = bold


def add_text(cell, text, bold=False, after=6, size=10.5):
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.16
    run = paragraph.add_run(text)
    set_run(run, bold, size)
    return paragraph


def main():
    document = Document(SOURCE)
    document.paragraphs[0].text = '深圳技术大学项目实践周记'
    document.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in document.paragraphs[0].runs:
        run.font.name = '黑体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        run.font.size = Pt(16)
        run.bold = True
    document.paragraphs[1].text = '第3周周记：2026年9月21日~2026年9月27日'
    document.paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in document.paragraphs[1].runs:
        set_run(run, size=10.5)

    table = document.tables[0]
    table.cell(2, 1).text = '多球种轻量追踪模型RK3588板端部署测试'
    for cell in table.rows[2].cells:
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.font.name = '宋体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                run.font.size = Pt(10.5)
                run.bold = True

    cell = table.cell(3, 0)
    cell.text = ''
    cell.vertical_alignment = 1
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run('一、工作内容')
    set_run(r, bold=True, size=11)

    add_text(cell, '本周配合算法同学钊沣完成LiteBallNet的RK3588板端验证。钊沣交接了网球、乒乓球、羽毛球三套模型权重、测试视频和PC参考脚本；我的工作是把这套PC流程转换并接入RK3588，输出与PC参考规则一致的追踪视频。')
    add_text(cell, '1.整理交接内容。先按钊沣的PC脚本核对模型输入、视频帧顺序、候选点处理和轨迹绘制方式，确定板端只替换模型推理部分，其余追踪和显示规则保持原样。')
    add_text(cell, '2.完成模型转换。使用PC虚拟机中的RKNN-Toolkit2.3.2，将三套PT模型导出ONNX并转换为RKNN。根据板端Runtime的输入要求，将运行器调整为NHWC布局传入9通道图像数据，避免每帧自动转置。正式效果测试选用FP模型，INT8版本仅保留作性能探索。')
    add_text(cell, '3.完成板端运行器。每套模型建立3个RKNN Runtime并分别绑定RK3588的NPU三核；板端输出沿用钊沣脚本中的候选点筛选、轨迹稳定、补点、异常点清理和红色拖尾绘制，仅在右上角增加24窗口滑动纯推理FPS。')
    add_text(cell, '4.完成视频与记录输出。三段视频均生成追踪视频、轨迹CSV和测试JSON，可直接用于与PC结果核对；乒乓球按原设定以30FPS慢放输出，网球和羽毛球保持原有输出策略。')

    add_text(cell, '二、测试结果', bold=True, size=11)
    result = cell.add_paragraph()
    result.paragraph_format.space_after = Pt(7)
    result.paragraph_format.line_spacing = 1.16
    for text, bold in [
        ('网球：279个三帧窗口，FP纯推理21.82FPS；', True),
        ('乒乓球：369个三帧窗口，FP纯推理22.60FPS；', True),
        ('羽毛球：182个三帧窗口，FP纯推理20.33FPS。', True),
        ('上述统计仅覆盖三路RKNN Runtime模型推理，排除视频解码、缩放、热图解码、轨迹绘制及视频写入。', False),
    ]:
        run = result.add_run(text)
        set_run(run, bold=bold)

    add_text(cell, '三、心得体会', bold=True, size=11)
    add_text(cell, '这次工作完成了三套模型从PC权重到板端视频输出的完整验证。后续可在更长视频和真实采集画面上继续检查遮挡、快速移动和切换场景时的稳定性。')

    for frame in FRAMES:
        picture = cell.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.paragraph_format.space_before = Pt(6)
        picture.paragraph_format.space_after = Pt(6)
        picture.add_run().add_picture(str(frame), width=Cm(15.2))

    sign = cell.add_paragraph()
    sign.paragraph_format.space_before = Pt(8)
    sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = sign.add_run('填写人：李睿鑫    2026年9月27日')
    set_run(run, size=10.5)
    document.save(OUTPUT)


if __name__ == '__main__':
    main()
