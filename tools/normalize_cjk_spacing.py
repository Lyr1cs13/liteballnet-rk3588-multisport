"""Remove accidental whitespace between CJK text and Latin letters or digits."""

from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DOCUMENT = Path(r"D:\文档\00项目与劳动实践\李睿鑫_周记3.docx")
CJK_TO_LATIN = re.compile(r"(?<=[\u3400-\u9fff])\s+(?=[A-Za-z0-9])")
LATIN_TO_CJK = re.compile(r"(?<=[A-Za-z0-9])\s+(?=[\u3400-\u9fff])")


def normalize(value: str) -> str:
    value = CJK_TO_LATIN.sub("", value)
    return LATIN_TO_CJK.sub("", value)


def disable_auto_spacing(paragraph) -> None:
    """Disable Word's visual auto-spacing between CJK, Latin text, and digits."""
    p_pr = paragraph._p.get_or_add_pPr()
    for tag in ("w:autoSpaceDE", "w:autoSpaceDN"):
        element = p_pr.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            p_pr.append(element)
        element.set(qn("w:val"), "0")


def paragraph_runs(document: Document):
    for paragraph in document.paragraphs:
        yield paragraph.runs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph.runs
    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                yield paragraph.runs
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            yield paragraph.runs


def paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        for container in (section.header, section.footer):
            yield from container.paragraphs
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs


def main():
    document = Document(DOCUMENT)
    for runs in paragraph_runs(document):
        for run in runs:
            # Assigning ``run.text`` clears embedded drawings, so leave image runs intact.
            if run._element.find(qn("w:drawing")) is None:
                run.text = normalize(run.text)
    for paragraph in paragraphs(document):
        disable_auto_spacing(paragraph)
    document.save(DOCUMENT)


if __name__ == "__main__":
    main()
